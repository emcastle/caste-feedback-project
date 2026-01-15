"""
document_docx.py

Purpose
-------
DOCX ingestion handler focused on extraction (not template parsing).

This module:
1) Reads a .docx using python-docx
2) Walks the document body in true order (paragraphs and tables interleaved)
3) Extracts ordered blocks:
   - Paragraphs
   - Table cells
4) Detects presence of equations (Office Math) and inserts placeholders
5) Produces two relational outputs:
   - documents_df: 1 row per DOCX (identity + metadata + doc-level errors)
   - blocks_df:    ordered blocks belonging to the document (for later parsing/segmentation)

Key Design Decisions
--------------------
- Extraction-only: no marker-based or template-based parsing here.
- Preserve ordering: block_num preserves true document order.
- Stable identity: doc_id is sha256 of DOCX bytes so it persists across runs.
- Equations: python-docx does not reliably extract equation text; we mark where
  equations occur with placeholders (and you can extend later to extract OMML XML).

Outputs (schemas)
-----------------
documents_df columns:
- doc_id
- source_file
- source_rel_path
- source_ext
- source_type
- num_blocks
- num_tables
- extractor_used
- error

blocks_df columns:
- doc_id
- block_num
- block_type          ("paragraph" | "table_cell")
- text
- table_index         (nullable)
- row_index           (nullable)
- col_index           (nullable)
- has_equation        (bool)
- error
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Iterator, List, Optional, Tuple

import pandas as pd


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        while True:
            b = f.read(chunk_size)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


def _clean_text(t: str) -> str:
    return (t or "").replace("\r", "\n").strip()


def _iter_block_items(doc) -> Iterator[object]:
    """
    Yield paragraphs and tables in document order.

    Based on python-docx internal patterns (walking doc.element.body children).
    """
    from docx.oxml.text.paragraph import CT_P  # type: ignore
    from docx.oxml.table import CT_Tbl  # type: ignore
    from docx.text.paragraph import Paragraph  # type: ignore
    from docx.table import Table  # type: ignore

    body = doc.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def _paragraph_has_equation(paragraph) -> bool:
    """
    Detect Office Math / equation runs via XML tags.
    OMML often appears as <m:oMath> or <m:oMathPara>.
    """
    xml = paragraph._p.xml  # underlying oxml string
    return ("m:oMath" in xml) or ("m:oMathPara" in xml)


def extract_docx_to_relational(
    docx_path: Path,
    source_rel_path: str,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    from docx import Document  # python-docx

    docx_path = docx_path.resolve()
    doc_id = sha256_file(docx_path)

    source_file = docx_path.name
    source_ext = ".docx"
    source_type = "docx_document"
    extractor_used = "python-docx(true-order)"

    try:
        doc = Document(str(docx_path))

        blocks: List[dict] = []
        block_num = 0
        table_count = 0

        for item in _iter_block_items(doc):
            # Paragraph
            if hasattr(item, "text") and item.__class__.__name__ == "Paragraph":
                txt = _clean_text(item.text)
                has_eq = _paragraph_has_equation(item)

                # If the paragraph is empty but contains an equation, preserve a placeholder
                if not txt and has_eq:
                    txt = "[EQUATION_PRESENT]"

                # If both text and equation exist, append marker (optional)
                elif txt and has_eq:
                    txt = txt + "\n[EQUATION_PRESENT]"

                if txt:
                    blocks.append(
                        {
                            "doc_id": doc_id,
                            "block_num": block_num,
                            "block_type": "paragraph",
                            "text": txt,
                            "table_index": None,
                            "row_index": None,
                            "col_index": None,
                            "has_equation": bool(has_eq),
                            "error": None,
                        }
                    )
                    block_num += 1

            # Table
            elif item.__class__.__name__ == "Table":
                t_i = table_count
                table_count += 1

                for r_i, row in enumerate(item.rows):
                    for c_i, cell in enumerate(row.cells):
                        cell_txt = _clean_text(cell.text)
                        if cell_txt:
                            blocks.append(
                                {
                                    "doc_id": doc_id,
                                    "block_num": block_num,
                                    "block_type": "table_cell",
                                    "text": cell_txt,
                                    "table_index": t_i,
                                    "row_index": r_i,
                                    "col_index": c_i,
                                    "has_equation": False,
                                    "error": None,
                                }
                            )
                            block_num += 1

        blocks_df = pd.DataFrame(blocks)

        documents_df = pd.DataFrame(
            [
                {
                    "doc_id": doc_id,
                    "source_file": source_file,
                    "source_rel_path": source_rel_path,
                    "source_ext": source_ext,
                    "source_type": source_type,
                    "num_blocks": int(len(blocks_df)),
                    "num_tables": int(table_count),
                    "extractor_used": extractor_used,
                    "error": None,
                }
            ]
        )

        return documents_df, blocks_df

    except Exception as e:
        documents_df = pd.DataFrame(
            [
                {
                    "doc_id": doc_id,
                    "source_file": source_file,
                    "source_rel_path": source_rel_path,
                    "source_ext": source_ext,
                    "source_type": source_type,
                    "num_blocks": None,
                    "num_tables": None,
                    "extractor_used": extractor_used,
                    "error": f"{type(e).__name__}: {e}",
                }
            ]
        )

        blocks_df = pd.DataFrame(
            [
                {
                    "doc_id": doc_id,
                    "block_num": None,
                    "block_type": None,
                    "text": "",
                    "table_index": None,
                    "row_index": None,
                    "col_index": None,
                    "has_equation": False,
                    "error": f"{type(e).__name__}: {e}",
                }
            ]
        )

        return documents_df, blocks_df


def build_full_text_from_blocks(blocks_df: pd.DataFrame, include_block_delimiters: bool = True) -> pd.DataFrame:
    """
    Convenience helper: derive doc-level full_text from blocks_df.

    Returns:
    - doc_id
    - full_text
    """
    if blocks_df.empty:
        return pd.DataFrame(columns=["doc_id", "full_text"])

    required = {"doc_id", "block_num", "text"}
    missing = required - set(blocks_df.columns)
    if missing:
        raise ValueError(f"blocks_df missing required columns: {sorted(missing)}")

    rows = []
    for doc_id, grp in blocks_df.groupby("doc_id", sort=False):
        g = grp.sort_values("block_num")
        if include_block_delimiters:
            parts = [
                f"=== BLOCK {int(r.block_num)} ({r.block_type}) ===\n{str(r.text).strip()}".strip()
                for r in g.itertuples(index=False)
            ]
        else:
            parts = [str(r.text).strip() for r in g.itertuples(index=False)]
        full_text = "\n\n".join([p for p in parts if p])
        rows.append({"doc_id": doc_id, "full_text": full_text})

    return pd.DataFrame(rows)